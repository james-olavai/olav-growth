#!/usr/bin/env python3
"""Red-team / self-check suite for extract_docs.py.

This file exists to try to falsify the new attack surface extract_docs.py
adds — sha256 dedup, markitdown/openpyxl parsing of untrusted office files,
the `--force` overwrite guard — not to raise a coverage number. Every test
below drives the real CLI in a subprocess against a throwaway DOCS_KB_ROOT,
so "the run survives a malformed file" is asserted as an actual process
exit code, not an in-process try/except that could differ from the CLI path.

Run (from this directory, with the docs-kb dependencies installed):
    python -m unittest redteam_extract_docs.py -v

Standard library only on purpose: this skill ships no test-framework
dependency, so this must not add one. openpyxl is the same dependency the
script under test already requires.
"""
import csv
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

SCRIPTS_DIR = Path(__file__).resolve().parent
EXTRACT = SCRIPTS_DIR / "extract_docs.py"


def run_extract(root, *args, set_root=True, timeout=180):
    """Run the real CLI; DOCS_KB_ROOT points at a throwaway dir unless a test
    deliberately clears it to exercise the ROOT_IS_DEFAULT refuse guard."""
    env = dict(os.environ)
    if set_root:
        env["DOCS_KB_ROOT"] = str(root)
    else:
        env.pop("DOCS_KB_ROOT", None)
    return subprocess.run(
        [sys.executable, str(EXTRACT), *args],
        cwd=str(SCRIPTS_DIR),
        env=env,
        capture_output=True,
        text=True,
        timeout=timeout,
    )


def garbage_bytes(name):
    """Deterministic bytes that magika won't classify as plain text. Pure
    ASCII garbage got routed to a text converter and "succeeded", so it would
    not exercise the format-specific parsers at all (and os.urandom made that
    outcome nondeterministic across runs)."""
    return bytes(range(256)) * 16 + name.encode("ascii")


def write_xlsx(path, sheets):
    """sheets: iterable of (sheet_name, list_of_rows)."""
    wb = openpyxl.Workbook()
    for index, (name, rows) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = name
        for row in rows:
            ws.append(row)
    wb.save(path)


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return [row for row in csv.reader(fh)]


class TempRootCase(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="rt-extract-docs-"))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)


class RootGuardTest(TempRootCase):
    def test_refuses_to_run_without_docs_kb_root(self):
        # The reused ROOT_IS_DEFAULT guard: without DOCS_KB_ROOT, ROOT would fall
        # back to this skill's own repo root. Must exit before touching disk.
        p = run_extract(self.root, set_root=False)
        self.assertEqual(p.returncode, 1, p.stderr)
        self.assertIn("Refusing to run", p.stderr)
        self.assertEqual(list(self.root.iterdir()), [])


class MalformedInputTest(TempRootCase):
    def test_binary_garbage_counted_failed_and_process_survives(self):
        names = ["g.docx", "g.pptx", "g.xlsx", "g.pdf"]
        for name in names:
            (self.root / name).write_bytes(garbage_bytes(name))
        write_xlsx(self.root / "good.xlsx", [("S", [["keep", "me"], [1, 2]])])

        p = run_extract(self.root)
        self.assertEqual(p.returncode, 0, f"run crashed:\n{p.stderr}")
        for name in names:
            self.assertIn(f"WARNING: extraction failed for {name}", p.stderr)
        self.assertIn("4 failed", p.stdout)
        # A valid sibling must still be fully processed in the same run.
        self.assertTrue((self.root / "good.xlsx.md").is_file())
        self.assertEqual(read_csv(self.root / "good.xlsx.S.csv"), [["keep", "me"], ["1", "2"]])
        # The atomic writer must not leave temp litter behind on failures.
        self.assertEqual([q.name for q in self.root.iterdir() if ".tmp" in q.name], [])

    def test_scanned_pdf_handoff_failure_is_contained(self):
        # < MIN_EXTRACTED_CHARS, so extract_docs hands the PDF to
        # ocr_extract.ocr_pdf(); pypdfium2 then rejects it. The point is that
        # the delegated-OCR failure stays inside the per-file try/except
        # instead of taking the whole run down.
        (self.root / "short.pdf").write_bytes(b"%PDF-1.4\n")
        write_xlsx(self.root / "ok.xlsx", [("S", [["x"]])])

        p = run_extract(self.root)
        self.assertEqual(p.returncode, 0, f"run crashed:\n{p.stderr}")
        self.assertIn("WARNING: extraction failed for short.pdf", p.stderr)
        self.assertIn("1 failed", p.stdout)
        self.assertTrue((self.root / "ok.xlsx.md").is_file())


class XlsxTrimTest(TempRootCase):
    def test_interior_blanks_kept_trailing_blanks_trimmed(self):
        # Data at rows 1-3, nothing until row 51 (col C only), plus a trailing
        # all-empty row 52. last_row/last_col must track the real extent, not
        # stop at the first gap and not keep the empty tail.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["h1", "h2"])
        ws.append([1, 2])
        ws.append([3, 4])
        ws.cell(row=51, column=3, value="tail")
        ws.cell(row=52, column=1, value=None)
        wb.save(self.root / "t.xlsx")

        p = run_extract(self.root)
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = read_csv(self.root / "t.xlsx.Data.csv")
        self.assertEqual(len(rows), 51, "interior gap/trailing row trimmed wrong")
        self.assertEqual(rows[0], ["h1", "h2", ""])       # truncated to last_col
        self.assertEqual(rows[2], ["3", "4", ""])
        self.assertEqual(rows[3], ["", "", ""])           # blank row kept as a gap
        self.assertEqual(rows[50], ["", "", "tail"])

    def test_sheet_name_sanitization_collision_keeps_both_sheets(self):
        # 'Q1 2026' and 'Q1_2026' both sanitize to the same CSV filename. The
        # second write used to silently replace the first sheet's export while
        # the index .md still advertised both sheets -> data loss.
        write_xlsx(
            self.root / "report.xlsx",
            [("Q1 2026", [["sheet-one", 1]]), ("Q1_2026", [["sheet-two", 2]])],
        )

        p = run_extract(self.root)
        self.assertEqual(p.returncode, 0, p.stderr)
        csvs = sorted(self.root.glob("report.xlsx.*.csv"))
        self.assertEqual(len(csvs), 2, [c.name for c in csvs])
        self.assertEqual(
            sorted(read_csv(c) for c in csvs),
            [[["sheet-one", "1"]], [["sheet-two", "2"]]],
        )
        index = (self.root / "report.xlsx.md").read_text(encoding="utf-8")
        self.assertIn("| Q1 2026 |", index)
        self.assertIn("| Q1_2026 |", index)


class XlsxResourceTest(TempRootCase):
    def test_many_mostly_empty_rows_complete_and_trim(self):
        # Bounded stand-in for "a hostile xlsx can't make extraction run away":
        # 20k rows x 5 cols with only two real cells must finish quickly and
        # export exactly the declared extent.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws.cell(row=1, column=1, value="header")
        ws.cell(row=20000, column=5, value="last")
        wb.save(self.root / "wide.xlsx")

        p = run_extract(self.root, timeout=120)
        self.assertEqual(p.returncode, 0, p.stderr)
        rows = read_csv(self.root / "wide.xlsx.S.csv")
        self.assertEqual(len(rows), 20000)
        self.assertEqual(rows[0], ["header", "", "", "", ""])
        self.assertEqual(rows[-1], ["", "", "", "", "last"])


class Sha256DedupTest(TempRootCase):
    def test_identical_different_names_stub_then_idempotent(self):
        write_xlsx(self.root / "a.xlsx", [("S", [["same", "content"]])])
        shutil.copyfile(self.root / "a.xlsx", self.root / "b.xlsx")

        first = run_extract(self.root)
        self.assertEqual(first.returncode, 0, first.stderr)
        self.assertIn("1 exact duplicates stubbed", first.stdout)
        stub = (self.root / "b.xlsx.md").read_text(encoding="utf-8")
        self.assertIn("extracted_via: sha256-duplicate", stub)
        self.assertIn("duplicate_of: a.xlsx", stub)
        # Only the canonical side gets a real export.
        self.assertTrue((self.root / "a.xlsx.md").is_file())
        self.assertEqual(
            sorted(q.name for q in self.root.glob("*.csv")), ["a.xlsx.S.csv"]
        )

        before = {q.name: q.read_bytes() for q in self.root.iterdir()}
        second = run_extract(self.root)
        self.assertEqual(second.returncode, 0, second.stderr)
        self.assertIn("0 failed", second.stdout)
        self.assertIn("0 exact duplicates stubbed", second.stdout)
        after = {q.name: q.read_bytes() for q in self.root.iterdir()}
        self.assertEqual(before, after, "second run rewrote or added files")


class ForceGuardTest(TempRootCase):
    def test_force_refuses_foreign_companion(self):
        write_xlsx(self.root / "foo.xlsx", [("S", [["real"]])])
        handwritten = self.root / "foo.xlsx.md"
        handwritten.write_text("# Handwritten notes\n\nDo not clobber me.\n", encoding="utf-8")
        original = handwritten.read_bytes()

        p = run_extract(self.root, "--force")
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertIn("WARNING: refusing to overwrite foo.xlsx.md", p.stderr)
        self.assertIn("1 failed", p.stdout)
        self.assertEqual(handwritten.read_bytes(), original)

    def test_force_overwrites_its_own_output(self):
        write_xlsx(self.root / "bar.xlsx", [("S", [["v1"]])])
        first = run_extract(self.root)
        self.assertEqual(first.returncode, 0, first.stderr)
        generated = self.root / "bar.xlsx.md"
        self.assertIn("extracted_source: true", generated.read_text(encoding="utf-8"))

        second = run_extract(self.root, "--force")
        self.assertEqual(second.returncode, 0, second.stderr)
        self.assertIn("1 succeeded", second.stdout)
        self.assertIn("extracted_source: true", generated.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
